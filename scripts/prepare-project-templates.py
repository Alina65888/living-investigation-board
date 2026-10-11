"""Editable estimate/register layouts measured from the supplied PDF references.
Run with the document runtime's Python. Output uses stored ZIP for browser filling.
No reference parties, amounts or signatures are embedded in the templates.
"""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_STORED
from lxml import etree
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / 'docs' / 'templates'

def attr(tag, **values):
    el = OxmlElement('w:'+tag)
    for k,v in values.items(): el.set(qn('w:'+k),str(v))
    return el

def paragraph(p,text='',bold=False,align=A.LEFT,keep=False):
    p.alignment=align
    f=p.paragraph_format
    f.space_before=f.space_after=Pt(0);f.line_spacing=1
    f.keep_with_next=keep;f.widow_control=False
    r=p.add_run(text);r.bold=bold;r.font.name='Times New Roman';r.font.size=Pt(12)
    return p

def table(doc,widths,borders=True):
    t=doc.add_table(rows=0,cols=len(widths));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,width in zip(t.columns,widths):col.width=Pt(width)
    pr=t._tbl.tblPr
    pr.append(attr('tblW',w=round(sum(widths)*20),type='dxa'))
    mar=attr('tblCellMar')
    for k,v in [('top',20),('bottom',20),('left',100),('right',100)]:mar.append(attr(k,w=v,type='dxa'))
    pr.append(mar)
    b=attr('tblBorders')
    for k in ['top','left','bottom','right','insideH','insideV']:b.append(attr(k,val='single' if borders else 'nil',sz=4,color='000000'))
    pr.append(b)
    return t

def row(t,texts,bold=False,aligns=None,header=False,keep=False):
    r=t.add_row()
    for c,text,i in zip(r.cells,texts,range(len(texts))):
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph(c.paragraphs[0],text,bold,(aligns or [A.LEFT]*len(texts))[i],keep)
    if header:r._tr.get_or_add_trPr().append(attr('tblHeader'))
    return r

def merged(t,label,total=None,bold=True,section=False,subtotal=False):
    r=t.add_row()
    if section:
        paragraph(r.cells[0].paragraphs[0],'{{sectionNumber}}',bold,A.CENTER,True)
        c=r.cells[1].merge(r.cells[-1]);paragraph(c.paragraphs[0],label,bold,A.LEFT,True)
    else:
        c=r.cells[1 if subtotal else 0].merge(r.cells[-2]);paragraph(c.paragraphs[0],label,bold,A.RIGHT if subtotal else A.LEFT)
        paragraph(r.cells[-1].paragraphs[0],total,bold,A.RIGHT)
    for c in r.cells:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return r

def mark(r,name):
    r._tr.addprevious(etree.Comment(name+':start'));r._tr.addnext(etree.Comment(name+':end'))

def build(kind):
    d=Document();s=d.sections[0]
    s.page_width=Pt(595.28);s.page_height=Pt(841.89)
    s.top_margin=s.bottom_margin=Pt(36);s.left_margin=Pt(36);s.right_margin=Pt(37)
    s.header_distance=s.footer_distance=Pt(18)
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1
    d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title=''
    compat=d.settings.element.find(qn('w:compat'))
    if compat is not None:compat.append(attr('doNotExpandShiftReturn'))
    if kind!='competition':
        paragraph(d.add_paragraph(),f"Приложение № {'2' if kind=='register' else '1'}\nк соглашению № {{{{agreementNumber}}}}\nот {{{{agreementDate}}}}",align=A.LEFT,keep=True).paragraph_format.left_indent=Pt(350)
        paragraph(d.add_paragraph(),'').paragraph_format.space_after=Pt(0)
    if kind=='register':
        paragraph(d.add_paragraph(),'Реестр документов',align=A.CENTER,keep=True)
        paragraph(d.add_paragraph(),'{{registryTitle}}',align=A.CENTER,keep=True)
        widths=[24,180,99,151,69]
        headers=['№','Наименование расходов','Контрагент','Перечень документов','Сумма (руб.)']
    else:
        paragraph(d.add_paragraph(),'{{heading}}',True,A.CENTER,True)
        if kind=='subsidy':paragraph(d.add_paragraph(),'{{subtitle}}',True,A.CENTER,True)
        widths=[41,221,58,53,73,77] if kind=='subsidy' else [27,233,54,59,77,72]
        headers=['№\nп/п' if kind=='subsidy' else '№','Наименование расходов' if kind=='subsidy' else 'Наименование','Ед. изм.','Кол-во','Стоимость\nза ед. (руб.)','Сумма\n(руб.)']
    if kind=='competition':d.paragraphs[-1].paragraph_format.space_after=Pt(10)
    else:paragraph(d.add_paragraph(),'').paragraph_format.space_after=Pt(0)
    t=table(d,widths)
    row(t,headers,True,[A.CENTER]*len(headers),True,True)
    if kind=='register':
        r=row(t,['{{rowNumber}}','{{rowName}}','{{rowParty}}','{{rowDocuments}}','{{rowTotal}}'],aligns=[A.CENTER,A.LEFT,A.LEFT,A.LEFT,A.RIGHT]);mark(r,'registry')
        r._tr.getprevious().addprevious(etree.Comment('body:start'))
        r._tr.getnext().addnext(etree.Comment('body:end'))
        merged(t,'ИТОГО:','{{total}}')
    else:
        first=merged(t,'{{sectionName}}',section=True);mark(first,'section')
        r=row(t,['{{itemNumber}}','{{itemText}}','{{itemUnit}}','{{itemQuantity}}','{{itemPrice}}','{{itemTotal}}'],aligns=[A.CENTER,A.JUSTIFY if kind=='subsidy' else A.LEFT,A.CENTER,A.CENTER,A.RIGHT,A.RIGHT]);mark(r,'item')
        last=merged(t,'ВСЕГО:' if kind=='subsidy' else 'ИТОГО:','{{subtotal}}',subtotal=True);mark(last,'subtotal')
        first._tr.getprevious().addprevious(etree.Comment('body:start'))
        last._tr.getnext().addnext(etree.Comment('body:end'))
        merged(t,'Общий ИТОГ:','{{total}}')
    for n,r in enumerate(t.rows):
        r._tr.get_or_add_trPr().append(attr('trHeight',val=780 if n==0 else 500,hRule='atLeast'))
    if kind=='subsidy':
        r=t.rows[2]
        for c in r.cells[:2]:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    paragraph(d.add_paragraph(),'').paragraph_format.space_after=Pt(12)
    sig=table(d,[350,172],False)
    row(sig,['{{signerPosition}}','{{signerName}}'],aligns=[A.LEFT,A.RIGHT])
    if kind=='subsidy':
        p=paragraph(s.footer.paragraphs[0],align=A.CENTER)
        r=p.add_run();field=attr('fldSimple',instr='PAGE');field.append(attr('r'));r._r.addnext(field)
    memory=BytesIO();d.save(memory)
    name='project-register' if kind=='register' else 'estimate-'+kind
    with ZipFile(memory) as src,ZipFile(OUT/(name+'.docx'),'w',compression=ZIP_STORED) as dest:
        for n in src.namelist():
            content=src.read(n)
            if n in ['word/styles.xml','word/stylesWithEffects.xml']:
                root=etree.fromstring(content)
                for style in list(root):
                    if style.tag==qn('w:style') and style.get(qn('w:styleId')) not in ['Normal','DefaultParagraphFont','TableNormal','NoList']:root.remove(style)
                content=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
            dest.writestr(n,content)
    print(name)

if __name__=='__main__':
    for kind in ['competition','subsidy','register']:build(kind)
